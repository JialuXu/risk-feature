"""导出最终交付物: 衍生指标设计稿.md / .csv / 加工需求文档.docx."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any

from indicator_pipeline.io_utils import write_csv

logger = logging.getLogger(__name__)


def export_md(out_dir: Path, batch_id: str, proposals: list[dict[str, Any]]) -> Path:
    """衍生指标设计稿 markdown."""
    L = []
    W = L.append
    W(f"# 衍生指标设计稿 · batch={batch_id}")
    W("")
    W(f"生成时间: {datetime.now().isoformat(timespec='seconds')}")
    W(f"指标条数: {len(proposals)}")
    W("")
    by_domain: dict[str, list[dict]] = {}
    for p in proposals:
        d = p.get("domain", "?")
        by_domain.setdefault(d, []).append(p)
    W("## 概览")
    W("")
    W("| domain | 条数 |")
    W("|---|---|")
    for k, v in by_domain.items():
        W(f"| {k} | {len(v)} |")
    W("")

    section_no = 2
    for domain, plist in by_domain.items():
        W(f"## {section_no}. {domain} 域 ({len(plist)} 条)")
        W("")
        for p in plist:
            W(f"### `{p.get('ind_code')}` — {p.get('ind_name_cn')}")
            W("")
            W(f"- **优先级**: {p.get('priority')}")
            W(f"- **粒度/窗口**: {p.get('granularity')} / {p.get('window')}")
            W(f"- **业务口径**: {p.get('biz_definition')}")
            W(f"- **计算逻辑**:")
            W(f"  ```")
            for line in str(p.get("calc_logic", "")).split("\n"):
                W(f"  {line}")
            W(f"  ```")
            W(f"- **依赖基础表**: {p.get('source_tables')}")
            W(f"- **依赖字段**: {p.get('source_fields')}")
            W(f"- **参考日期口径**: {p.get('ref_date_logic')}")
            W(f"- **空值处理**: {p.get('null_handling')}")
            W(f"- **后处理**: {p.get('post_processing_rule')}")
            iv = p.get("current_iv")
            if iv is not None:
                W(f"- **当前 IV**: {iv} ({p.get('current_iv_credibility')}) · "
                  f"覆盖率: {p.get('current_coverage_rate')} · "
                  f"风险方向: {p.get('current_risk_direction')}")
            shadow = p.get("shadow_iv_status")
            if shadow:
                W(f"- **影子 IV 状态**: {shadow}")
            W("")
        section_no += 1

    out_path = out_dir / "衍生指标设计稿.md"
    out_path.write_text("\n".join(L), encoding="utf-8")
    logger.info("[delivery] 导出 md: %s", out_path)
    return out_path


def export_csv(out_dir: Path, proposals: list[dict[str, Any]]) -> Path:
    out_path = out_dir / "衍生指标设计稿.csv"
    cols = [
        "ind_code", "ind_name_cn", "domain", "priority",
        "granularity", "window", "biz_definition", "calc_logic",
        "source_tables", "source_fields",
        "current_iv", "current_iv_credibility", "current_coverage_rate",
        "current_risk_direction", "ref_date_logic", "null_handling",
        "post_processing_rule", "shadow_iv_status", "shadow_iv",
        "lifecycle_status", "ind_version",
    ]
    rows = []
    for p in proposals:
        r = {}
        for c in cols:
            v = p.get(c)
            if isinstance(v, (list, dict)):
                import json
                v = json.dumps(v, ensure_ascii=False)
            r[c] = v
        rows.append(r)
    write_csv(out_path, rows, fieldnames=cols)
    logger.info("[delivery] 导出 csv: %s", out_path)
    return out_path


def export_docx(out_dir: Path, batch_id: str, proposals: list[dict[str, Any]]) -> Path | None:
    """加工需求文档 docx (给数仓/数开). 失败则降级跳过."""
    try:
        from docx import Document
    except ImportError:
        logger.warning("[delivery] 缺少 python-docx, 跳过 docx 导出")
        return None

    doc = Document()
    doc.add_heading(f"衍生指标加工需求 · batch={batch_id}", 0)
    doc.add_paragraph(f"生成时间: {datetime.now().isoformat(timespec='seconds')}")
    doc.add_paragraph(f"指标条数: {len(proposals)}")

    # 按 domain 分章节
    by_domain: dict[str, list[dict]] = {}
    for p in proposals:
        by_domain.setdefault(p.get("domain", "?"), []).append(p)

    for domain, plist in by_domain.items():
        doc.add_heading(f"{domain} 域 ({len(plist)} 条)", level=1)
        for p in plist:
            doc.add_heading(f"{p.get('ind_code')} — {p.get('ind_name_cn')}", level=2)
            t = doc.add_table(rows=0, cols=2)
            t.style = "Light Grid Accent 1"
            for k_label, v_key in [
                ("优先级", "priority"),
                ("业务口径", "biz_definition"),
                ("计算逻辑", "calc_logic"),
                ("依赖基础表", "source_tables"),
                ("依赖字段", "source_fields"),
                ("参考日期口径", "ref_date_logic"),
                ("空值处理", "null_handling"),
                ("后处理", "post_processing_rule"),
                ("当前 IV", "current_iv"),
            ]:
                row = t.add_row().cells
                row[0].text = k_label
                v = p.get(v_key)
                if isinstance(v, (list, dict)):
                    import json
                    v = json.dumps(v, ensure_ascii=False)
                row[1].text = str(v) if v is not None else ""
            doc.add_paragraph("")

    out_path = out_dir / "加工需求文档.docx"
    doc.save(str(out_path))
    logger.info("[delivery] 导出 docx: %s", out_path)
    return out_path
